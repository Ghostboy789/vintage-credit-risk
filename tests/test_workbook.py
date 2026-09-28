"""Recalculates excel/vintage_ecl_workbook.xlsx with the `formulas` package and checks its
live-formula totals against artefacts/ecl.json, within rounding.
"""

import json
from pathlib import Path

import formulas
import pytest

REPO_ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = REPO_ROOT / "excel" / "vintage_ecl_workbook.xlsx"
ARTEFACTS = REPO_ROOT / "artefacts"
LATEST_DATE = "2026-03-01"

RTOL = 1e-6  # "within rounding": openpyxl/formulas carry full float precision, so this is tight


def expected_totals():
    ecl = json.load(open(ARTEFACTS / "ecl.json", encoding="utf-8"))
    rows = [r for r in ecl["by_date"] if r["reporting_date"] == LATEST_DATE]
    total_ead = sum(r["ead"]["value"] for r in rows)
    ecl_12m = sum(r["ecl"]["value"] for r in rows if r["stage"] == "1")
    ecl_lifetime = sum(r["ecl"]["value"] for r in rows if r["stage"] in ("2", "3"))
    return {
        "total_ead": total_ead,
        "ecl_12m": ecl_12m,
        "ecl_lifetime": ecl_lifetime,
        "total_ecl": ecl_12m + ecl_lifetime,
        "coverage": (ecl_12m + ecl_lifetime) / total_ead,
    }


def test_recalculated_totals_match_ecl_json():
    """Recalculate the whole workbook and check every headline ECL total against ecl.json."""
    xl = formulas.ExcelModel().loads(str(WORKBOOK)).finish()
    sol = xl.calculate()
    exp = expected_totals()

    def find(name):
        # `formulas` keys defined names as '[FILE]SHEET'!NAME (case-insensitive on file/sheet)
        for key, cell in sol.items():
            if key.upper().endswith(f"'!{name}"):
                return cell.value[0, 0] if hasattr(cell.value, "shape") else cell.value
        raise KeyError(name)

    assert find("TOTAL_EAD") == pytest.approx(exp["total_ead"], rel=RTOL)
    assert find("ECL_12M") == pytest.approx(exp["ecl_12m"], rel=RTOL)
    assert find("ECL_LIFETIME") == pytest.approx(exp["ecl_lifetime"], rel=RTOL)
    assert find("TOTAL_ECL") == pytest.approx(exp["total_ecl"], rel=RTOL)


def test_workbook_has_no_small_cells():
    """Publication rule: no ECL-sheet row may describe fewer than 10 loans."""
    from openpyxl import load_workbook

    wb = load_workbook(WORKBOOK, data_only=False)
    ws = wb["ECL"]
    # The grade x stage table starts at row 5 (below the header row) and runs for exactly
    # 7 grades x 3 stages = 21 rows; rows after that are portfolio totals and the sensitivity
    # grid, which don't carry an n_loans column.
    for row in ws.iter_rows(min_row=5, max_row=25, max_col=3):
        n_loans = row[2].value
        if isinstance(n_loans, (int, float)):
            assert n_loans >= 10, f"row {row[0].row} describes fewer than 10 loans"
